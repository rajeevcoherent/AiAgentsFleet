"""PLAN-08: client-ready document and a blocking release gate.

The gate (audit.release_gate) turns any blocker into ready = false; the client
appendix carries only A.1-A.3; QA year attribution, value-chain ranges and
register IDs are fixed. No web search and no live LLM: the plausibility review
runs against a stub."""
from pathlib import Path

import docx
import pytest

from AiAgentsFleet.cmi_report import audit, build, factcheck, model
from AiAgentsFleet.cmi_report.register import build_register
from AiAgentsFleet.cmi_report.run_cmi_report import clear_charts, draft_path, finalize

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "templates" / "cmi_whitepaper_template.docx"
N = model.node


def _m(**over) -> dict:
    """A small but complete model: two regions, MEA sub-regions, a few countries and segments."""
    m = {
        "market": "Global Widget Market",
        "global": {**N(100, 200), "v25": 92, "series": {2025: 92, 2026: 100, 2033: 200}},
        "regions": {"North America": {**N(60, 110), "v25": 56, "series": {2025: 56, 2026: 60, 2033: 110}},
                    "Middle East & Africa": {**N(40, 90), "v25": 36, "series": {2025: 36, 2026: 40, 2033: 90}}},
        "subregions": {"Middle East": {**N(25, 55), "v25": 23, "series": {}},
                       "Africa": {**N(15, 35), "v25": 13, "series": {}}},
        "countries": {"North America": {"United States": N(50, 90), "Canada": N(10, 20)},
                      "Middle East": {"Saudi Arabia": N(15, 35), "Rest of Middle East": N(10, 20)},
                      "Africa": {"Morocco": N(5, 15), "Rest of Africa": N(10, 20)}},
        "country_series": {},
        "segments": {"Global": {"type": {"Hardware": N(60, 110), "Software": N(40, 90)}, "demand": {}},
                     "Middle East": {"type": {"Hardware": N(15, 30), "Software": N(10, 25)}, "demand": {}},
                     "Africa": {"type": {"Hardware": N(9, 20), "Software": N(6, 15)}, "demand": {}},
                     "Morocco": {"type": {"Ascorbic Acid (pure)": N(2, 6), "Ascorbic Acid derivatives": N(3, 9)},
                                 "demand": {}}},
        "sections": {"type": {"show": True, "why": ""}, "demand": {"show": False, "why": "no data"},
                     "regulation": {"show": False, "why": ""}, "milestones": {"show": False, "why": ""},
                     "vendor_mea": {"show": True, "why": ""}},
        "dimensions": {"type": {"singular": "Type", "plural": "Types"},
                       "demand": {"singular": "End User", "plural": "End Users"}},
        "scenarios": {"Base": {"cagr": 10.4, "series": {2033: 200.0}}},
        "vendor_share_global": [("Acme", 12.0), ("Others", 88.0)], "vendor_share_mea": [],
        "vendor_chart_global": [("Acme", 12.0), ("Regional & other", 88.0)], "vendor_chart_mea": [],
        "value_chain": [], "published_2025": [],
        "profiled": {"Middle East": ["Saudi Arabia"], "Africa": ["Morocco"]},
        "definition": {"definition": "Widgets sold to end users."},
        "size_check": {"status": "sourced", "notes": [], "source": {"quote": "q", "publisher": "P", "domain": "p.com"}},
        "geo_review": {"blockers": [], "warnings": []},
        "segment_review": {"blockers": [], "warnings": []},
        "corrections": [], "allocation_rules": [],
    }
    m.update(over)
    return m


def _doc(tmp_path: Path, extra: list[tuple[str, str]] = (), key_items: int = 3, profile_lead: bool = True,
         name: str = "doc.docx") -> Path:
    """A finished-report stand-in: headings with bodies, a Key conclusions box, country profiles."""
    d = docx.Document()
    d.add_heading("01   Executive Summary", 1)
    d.add_paragraph("CMI estimates the global widget market at $100M in 2026E, rising to $200M by 2033F.")
    t = d.add_table(rows=1, cols=1)
    cell = t.rows[0].cells[0]
    cell.paragraphs[0].text = "Key conclusions"
    for k in range(key_items):
        cell.add_paragraph(f"{k + 1}. Conclusion number {k + 1} about the widget market.")
    d.add_heading("04   Middle East Widget Market", 1)
    d.add_heading("4.4 Middle East Widget Market: Country landscape", 2)
    d.add_heading("Saudi Arabia", 3)
    if profile_lead:
        d.add_paragraph("Saudi Arabia is a $15.0M market in 2026E (CMI estimate), the region's largest.")
    d.add_heading("Saudi Arabia Opportunity by Type", 3)
    d.add_paragraph("Hardware leads the Saudi market throughout the forecast period.")
    d.add_heading("05   Africa Widget Market", 1)
    d.add_heading("Morocco", 3)
    d.add_paragraph("Morocco is a $5.0M market in 2026E (CMI estimate), growing quickly from a small base.")
    for style, text in extra:
        if style.startswith("Heading"):
            d.add_heading(text, int(style[-1]))
        else:
            d.add_paragraph(text)
    d.add_heading("11   Appendix: Master Data and Registers", 1)
    d.add_heading("A.2 Number register (derived facts)", 2)
    d.add_paragraph("Each derived figure used in the report, with its register ID.")
    p = tmp_path / name
    d.save(str(p))
    return p


QA_CLEAN = {"numbers_not_in_model": [], "same_fact_inconsistencies": {}, "repeated_sentences": []}
PLAUS_CLEAN = {"ran": True, "items": []}


def _gate(m, path, qa=None, computed=None, plaus=PLAUS_CLEAN, fact_review=None):
    return audit.release_gate(m, audit.doc_blocks(path), qa_report=qa or QA_CLEAN, computed=computed or {},
                              plausibility=plaus, fact_review=fact_review)


# --------------------------------------------------------------------------- gate
def test_clean_doc_is_ready(tmp_path):
    g = _gate(_m(), _doc(tmp_path))
    assert g["ready"], g["blockers"]
    assert g["blockers"] == []


def test_to_verify_marker_is_not_ready(tmp_path):
    g = _gate(_m(), _doc(tmp_path, extra=[("", "Turkey | BUGEM [to verify] | Registration required.")]))
    assert not g["ready"]
    assert any("[to verify]" in b for b in g["blockers"])


@pytest.mark.parametrize("text", ["A note: TBD.", "Lorem ipsum dolor sit amet.", "Drones are used for spraying.",
                                  "Item [unverified] here."])
def test_other_leftovers_block(tmp_path, text):
    assert not _gate(_m(), _doc(tmp_path, extra=[("", text)]))["ready"]


@pytest.mark.parametrize("text", ["51 generated sentence(s) removed before publication.",
                                  "Number consistency sweep of the finished report: 900 numbers checked.",
                                  "A.5 Corrections made during the self-audit"])
def test_internal_wording_blocks(tmp_path, text):
    g = _gate(_m(), _doc(tmp_path, extra=[("", text)]))
    assert any(b.startswith("internal pipeline wording") for b in g["blockers"]), g["blockers"]


def test_register_id_wording_only_allowed_in_a2(tmp_path):
    assert _gate(_m(), _doc(tmp_path))["ready"]                      # A.2 says "register ID": fine
    g = _gate(_m(), _doc(tmp_path, extra=[("", "Each figure carries a register ID for audit.")]))
    assert any("register ID" in b and "outside A.2" in b for b in g["blockers"])


def test_qa_numbers_and_inconsistencies_block(tmp_path):
    p = _doc(tmp_path)
    assert not _gate(_m(), p, qa={**QA_CLEAN, "numbers_not_in_model": [{"text": "x"}]})["ready"]
    assert not _gate(_m(), p, qa={**QA_CLEAN, "same_fact_inconsistencies": {"Egypt value 2026": {}}})["ready"]
    assert not _gate(_m(), p, computed={"math_sums_cagr_increments": ["regions sum to 99 in 2026"]})["ready"]


def test_unsourced_headline_with_published_wording_blocks(tmp_path):
    p = _doc(tmp_path, extra=[("", "Published 2025 estimates vary between $90M and $120M; CMI triangulated them.")])
    assert _gate(_m(), p)["ready"]                                   # sourced headline: comparison allowed
    g = _gate(_m(size_check={"status": "unsourced", "notes": []}), p)
    assert any("unsourced" in b for b in g["blockers"])
    # an unsourced headline with no such wording is fine (it is labelled a CMI estimate)
    assert _gate(_m(size_check={"status": "unsourced", "notes": []}), _doc(tmp_path, name="b.docx"))["ready"]


@pytest.mark.parametrize("key", ["size_check", "geo_review", "segment_review", "vendor_review", "fact_review"])
def test_review_blockers_block(tmp_path, key):
    rev = {"status": "sourced", "blockers": ["something still failing after one re-ask"], "warnings": ["w"]}
    g = _gate(_m(**{key: rev}), _doc(tmp_path))
    assert not g["ready"]
    assert f"{key}: something still failing after one re-ask" in g["blockers"]
    assert f"{key}: w" in g["warnings"]


def test_absent_reviews_are_skipped_and_fact_checks_file_is_read(tmp_path):
    m = _m()
    for k in ("geo_review", "segment_review", "size_check"):
        m.pop(k)
    assert _gate(m, _doc(tmp_path))["ready"]
    g = _gate(m, _doc(tmp_path), fact_review={"blockers": ["Turkey: TAGEM is not the regulator"]})
    assert "fact_review: Turkey: TAGEM is not the regulator" in g["blockers"]


def test_unsourced_vendor_share_over_20_blocks(tmp_path):
    p = _doc(tmp_path)
    m = _m(vendor_share_global=[("DSM-Firmenich", 23.0), ("Others", 77.0)])
    g = _gate(m, p)
    assert any("DSM-Firmenich" in b for b in g["blockers"])
    m["vendor_sources"] = {"global": [{"vendor": "DSM-Firmenich", "share_pct": 23.0, "quote": "DSM holds 23%"}]}
    assert _gate(m, p)["ready"]


def test_empty_heading_key_conclusions_and_profile_lead(tmp_path):
    g = _gate(_m(), _doc(tmp_path, extra=[("Heading2", "9.1 Data triangulation"), ("Heading2", "9.2 Scope"),
                                          ("", "The scope covers widgets.")]))
    assert "empty heading: '9.1 Data triangulation'" in g["blockers"]
    g = _gate(_m(), _doc(tmp_path, key_items=1))
    assert any("Key conclusions" in b for b in g["blockers"])
    g = _gate(_m(), _doc(tmp_path, profile_lead=False))
    assert any("Saudi Arabia" in b and "lead" in b for b in g["blockers"])


def test_numbered_parent_heading_is_not_empty(tmp_path):
    p = _doc(tmp_path, extra=[("Heading2", "3.5 Regional Analysis, By Key Countries"),
                              ("Heading2", "North America Widget Market by Country"), ("", "The US leads.")])
    assert _gate(_m(), p)["ready"]


# --------------------------------------------------------------------------- plausibility review (stub LLM)
class StubLLM:
    def __init__(self, answer):
        self.answer, self.prompts = answer, []

    def complete_json(self, system, user, **k):
        self.prompts.append((system, user))
        if isinstance(self.answer, Exception):
            raise self.answer
        return self.answer


def test_plausibility_high_item_is_a_warning_not_a_blocker(tmp_path):
    llm = StubLLM({"items": [{"quote": "Alibaba Cloud at 15.4%", "issue": "More than the whole China market.",
                              "category": "vendor", "severity": "high", "fix": "Drop the share."},
                             {"quote": "x", "issue": "Wording.", "category": "logic", "severity": "low", "fix": ""}]})
    pr = audit.plausibility_review(_m(), "Report text with Alibaba Cloud at 15.4%.", llm=llm)
    assert pr["ran"] and len(pr["items"]) == 2
    system, user = llm.prompts[0]
    assert "sceptical" in system and "MASTER DATA SHEET" in user and "REPORT TEXT" in user
    assert "Generative AI in Agriculture" in user                     # FINDINGS sections B-F as few-shot
    g = _gate(_m(), _doc(tmp_path), plaus=pr)
    # user decision 2026-10-16: AI review items are warnings for the human reviewer, never blockers
    assert g["ready"]
    assert any(w.startswith("plausibility review (high) vendor") for w in g["warnings"])


def test_plausibility_medium_is_a_warning_and_failure_is_a_warning(tmp_path):
    p = _doc(tmp_path)
    pr = audit.plausibility_review(_m(), "text", llm=StubLLM({"items": [
        {"quote": "q", "issue": "Check this.", "category": "fact", "severity": "medium", "fix": "f"}]}))
    g = _gate(_m(), p, plaus=pr)
    assert g["ready"] and any("medium" in w for w in g["warnings"])
    g2 = _gate(_m(), p, plaus=audit.plausibility_review(_m(), "t", llm=StubLLM(RuntimeError("down"))))
    assert g2["ready"] and any("did not run" in w for w in g2["warnings"])
    assert _gate(_m(), p, plaus=None)["ready"]                       # --no-ai: a warning, not a blocker


# --------------------------------------------------------------------------- internal QA report
def test_internal_qa_report_lists_gate_result(tmp_path):
    m = _m()
    g = _gate(m, _doc(tmp_path, extra=[("", "Kenya | Communications Authority [to verify] | Data rules.")]))
    fleet = {"fields_raw": {"market_size": {"base_year_value_usd_mn": 90, "cagr_pct": 10.0},
                            "regional": {"regions": {"North America": {"value_usd_mn": 50,
                                                                       "countries": [{"name": "United States",
                                                                                      "value_usd_mn": 45}]}}}}}
    qa = {**QA_CLEAN, "internal_log": {"to_verify": ["Kenya | Communications Authority [to verify]"],
                                       "corrections": ["c1"], "consistency_log": ["l1"]}}
    p = audit.write_internal_qa(tmp_path / "x-internal-qa.md", m, fleet, gate=g, qa_report=qa,
                                audit_res={"plausibility_review": PLAUS_CLEAN, "computed": {}},
                                prose={"dropped": {"ch2": ["[repeat] a sentence"]}}, client_doc="x.docx")
    text = p.read_text(encoding="utf-8")
    assert "## 1. Release gate" in text and "**NOT READY**" in text and "[to verify]" in text
    for sec in ("Headline market-size check", "Geography review", "Fact checks", "Dropped sentences",
                "Raw vs final", "AI plausibility review"):
        assert sec in text
    assert "| Global 2025 (USD mn) | 90 | 92.0 | +2.2% |" in text


# --------------------------------------------------------------------------- DRAFT renaming, stale charts
def test_finalize_renames_to_draft_and_back(tmp_path):
    out = tmp_path / "x-cmi-whitepaper.docx"
    out.write_text("d")
    out.with_suffix(".pdf").write_text("p")
    assert finalize(out, False) == draft_path(out)
    assert not out.exists() and draft_path(out).exists() and (tmp_path / "x-cmi-whitepaper.DRAFT.pdf").exists()
    out.write_text("d2")
    assert finalize(out, True) == out
    assert out.exists() and not draft_path(out).exists() and not (tmp_path / "x-cmi-whitepaper.DRAFT.pdf").exists()


def test_clear_charts_removes_stale_pngs(tmp_path):
    c = tmp_path / "charts"
    c.mkdir()
    (c / "demand_qatar.png").write_bytes(b"x")
    (c / "notes.txt").write_text("keep")
    assert clear_charts(c) == 1
    assert not (c / "demand_qatar.png").exists() and (c / "notes.txt").exists()


# --------------------------------------------------------------------------- QA year attribution (FINDINGS D)
GENAI = ("Demand centres on Farm Advisory & Recommendations, the largest Application Type in 2026E at $0.35M and "
         "in 2033F at $1.18M, CMI forecasts, as water-scarcity pressures push growers toward advisory tools.")


def _reg(name, v26, v33):
    return [{"id": "S-1", "fact": f"{name} market size, 2026E", "entities": [name], "display": v26, "metric": "value"},
            {"id": "S-2", "fact": f"{name} market size, 2033F", "entities": [name], "display": v33, "metric": "value"}]


def test_genai_sentence_is_not_a_false_flag():
    reg = _reg("Farm Advisory & Recommendations", "$0.35M", "$1.18M") + _reg("Rest of Middle East", "$0.10M", "$0.35M")
    fc = factcheck.check(reg, [GENAI], [])
    assert fc["mismatches"] == [], fc["mismatches"]


@pytest.mark.parametrize("text,year", [("Egypt reaches $5.0M by 2033F.", "2033"),
                                       ("Egypt is $2.0M in 2026E.", "2026"),
                                       ("Egypt ($2.0M, 2026E) leads.", "2026"),
                                       ("In 2026E Egypt is at $2.0M and in 2033F at $5.0M.", None)])
def test_year_binding(text, year):
    reg = _reg("Egypt", "$2.0M", "$5.0M")
    assert factcheck.check(reg, [text], [])["mismatches"] == []
    if year:
        tok = "$5.0M" if year == "2033" else "$2.0M"
        i = text.index(tok)
        assert factcheck._year_near(text, i, i + len(tok)) == year


def test_wrong_year_value_is_still_flagged():
    fc = factcheck.check(_reg("Egypt", "$2.0M", "$5.0M"), ["Egypt reaches $2.0M by 2033F."], [])
    assert fc["mismatches"]


# --------------------------------------------------------------------------- value chain, register IDs
def test_value_chain_low_not_above_high():
    notes = []
    vc = model.value_chain_ranges([{"stage": "Raw materials", "share_low_pct": 20, "share_high_pct": 30},
                                   {"stage": "Vitamin C Production", "share_low_pct": 36, "share_high_pct": 35},
                                   {"stage": "Distribution", "share_low_pct": 40, "share_high_pct": 40},
                                   {"stage": "Retail", "share_low_pct": 30, "share_high_pct": 40}], notes)
    assert any("rescaled" in n for n in notes)
    for s in vc:
        assert s["share_low_pct"] <= s["share_high_pct"] and s["share_high_pct"] - s["share_low_pct"] >= 1, s
    assert any("Vitamin C Production" in n for n in notes)


def test_register_ids_unique_when_codes_collide():
    m = _m()
    reg = build_register(m, {"dynamics": {"mea_drivers": [{"name": "Food security", "impact_pp": 1.0}],
                                          "mea_restraints": [{"name": "Food security", "impact_pp": -0.5}]}})
    ids = [e["id"] for e in reg]
    assert len(ids) == len(set(ids))
    pure = {e["id"] for e in reg if e["fact"].startswith("Morocco Type: Ascorbic Acid (pure) market size, 2026")}
    der = {e["id"] for e in reg if e["fact"].startswith("Morocco Type: Ascorbic Acid derivatives market size, 2026")}
    assert pure and der and pure != der
    assert pure == {"S-MOROCCO-ASCORBICAC-2026"}                       # the first keeps the plain code


# --------------------------------------------------------------------------- client document
def _blocks(els):
    return [e.xpath("string(.)") for e in els]


def test_to_verify_items_skip_the_explanatory_sentence():
    d = docx.Document()
    d.add_paragraph("Shares are CMI assessments; items marked [to verify] could not be confirmed.")
    d.add_paragraph("Kenya's regulator is the PCPB [to verify].")
    b = type("B", (), {"out": [p._p for p in d.paragraphs]})()
    items = build.to_verify_items(b, {}, _m())
    assert items == ["Kenya's regulator is the PCPB [to verify]."]


def test_builder_drops_empty_headings():
    d = docx.Document()
    d.add_heading("09   Research Methodology", 1)
    d.add_paragraph("CMI's estimates combine regional and national totals.")
    d.add_heading("9.1 Data triangulation", 2)
    d.add_heading("9.2 Scope", 2)
    d.add_paragraph("Widgets.")
    d.add_heading("10   Empty chapter", 1)
    d.add_heading("10.1 Empty section", 2)
    d.add_heading("11   Appendix", 1)
    d.add_paragraph("Data.")
    b = build.Builder.__new__(build.Builder)
    b.out, b.log = [e for e in d.element.body.iterchildren() if not e.tag.endswith("sectPr")], []
    b.drop_empty_headings()
    texts = [t for t in _blocks(b.out) if t]
    assert "9.1 Data triangulation" not in texts and "10.1 Empty section" not in texts
    assert "10   Empty chapter" not in texts                            # left empty once its sub-section went
    assert "9.2 Scope" in texts and "09   Research Methodology" in texts
    assert sum(x.startswith("empty heading removed") for x in b.log) == 3




def test_triangulation_allowed_only_in_the_methodology_chapter(tmp_path):
    """The VP's methodology text names data triangulation (2026-10-19); with an unsourced headline the
    word still blocks anywhere else in the report."""
    m = _m(size_check={"status": "unsourced", "notes": []})
    ok = _doc(tmp_path, extra=[("Heading1", "09   Research Methodology"),
                               ("", "CMI combines secondary research and data triangulation.")], name="a.docx")
    assert _gate(m, ok)["ready"], _gate(m, ok)["blockers"]
    bad = _doc(tmp_path, extra=[("", "CMI triangulated the published estimates.")], name="b.docx")
    assert any("unsourced" in b for b in _gate(m, bad)["blockers"])
