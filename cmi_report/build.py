"""Fills the CMI whitepaper template (the drone-market .docx) with one
market's model, prose, agent tables and charts.

The template is cloned element by element (body indices below refer to
`Updated_Drone Market_Whitepaper Template_CMI.docx`); each element keeps its
styling and only its content changes. Figures and tables are renumbered in
order. Removed on purpose: the 'Data status' box, the Contact section (with
its email/phone placeholders) and the Disclaimer, plus the 'draft template'
wording in the footer and the assumptions list.
"""
from __future__ import annotations

import copy
import re
from datetime import datetime
from pathlib import Path

import docx
from docx.oxml.ns import qn
from PIL import Image

from AiAgentsFleet.cmi_report import fmt
from AiAgentsFleet.cmi_report.guard import ClaimGuard, NumberGuard, clean_text, vendor_clean
from AiAgentsFleet.cmi_report.model import Y0, Y1, regional_chapters
from AiAgentsFleet.cmi_report.guard import sentences
from AiAgentsFleet.cmi_report.groups import canonical, short
from AiAgentsFleet.cmi_report.prose import (absent_names, base_facts, claim_groups, fill_conclusions, method_text,
                                            highlight_regions, seg_insight, with_lead)

_COUNT = {1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five", 6: "Six"}


def _sentence(text: str) -> str:
    text = (text or "").strip()
    return text if not text or text[-1] in ".!?" else text + "."


_MID_THOUGHT = re.compile(r"^\s*(?:within|also|the region|this market|it|its|they|these|this)\b"
                          r"|^\W*(?:\w+\W+){0,5}?(?:also|too)\b", re.I)


def open_chapter(text: str | None, must: str, lead: str, avoid: str | None = None) -> str:
    """A chapter's first sentence names the chapter's own subject (`must`) and does not start
    mid-thought ('Within Africa, ...', 'X is also ...'); otherwise that sentence goes and a
    data-based opener (`lead`) comes first. `avoid`: a subject that must not lead this chapter
    (the other sub-region)."""
    sents = sentences(text or "")
    if not sents:
        return lead
    first = sents[0]
    # the other sub-region has no place in this chapter's opening sentence ('The Middle East accounts for
    # $18.4M and Africa $10.8M' opens the Middle East chapter on Africa's number)
    wrong_subject = bool(avoid) and bool(re.search(avoid, first))
    if re.search(must, first) and not _MID_THOUGHT.search(first) and not wrong_subject:
        return " ".join(sents)
    rest = sents[1:] if (_MID_THOUGHT.search(first) or wrong_subject) else sents
    return " ".join([lead] + rest)


def complete_rows(rows: list[list[str]], log: list[str], what: str) -> list[list[str]]:
    """Agent-written rows with an empty cell are left out rather than shown with '-'."""
    keep = [r for r in rows if all(str(c).strip() and str(c).strip() not in ("-", "–") for c in r)]
    if len(keep) < len(rows):
        log.append(f"{what}: {len(rows) - len(keep)} incomplete row(s) left out")
    return keep

DEFINITION_CH = 0          # the definition chapter's original number: it comes first
DEFINITION_TITLE = "Market Definition & Taxonomy"
_MEA_WORDS = re.compile(r"\b(?:Middle East|Africa|African|MEA|Gulf|GCC)\b")


def global_only(text: str) -> str:
    """The sentences of `text` that do not mention the Middle East & Africa (CEO feedback 2026-10-20: the
    summary's opening paragraph is about the global market only)."""
    return " ".join(x for x in sentences(text or "") if not _MEA_WORDS.search(x))


REGION_CH = 20            # original numbers of the regional chapters (the template's chapters are 1-10)
REMOVE = {22, 76, 77, 78, 79, 80, 393, 394, 395, 396}
# every country table: the last column names the opportunity and its period (user decision 2026-10-18)
COUNTRY_HEADER = ["Country", f"{Y0}E", f"{Y1}F", "CAGR", f"Incremental $ Opportunity ({Y0}-{str(Y1)[2:]})"]
PROFILE_BLOCKS = {  # first element of each country profile block (12 elements each)
    "Middle East": [169, 182, 195, 208],
    "Africa": [249, 261, 273, 285],
}
PROFILE_STARTS = {v[0]: k for k, v in PROFILE_BLOCKS.items()}           # 169 -> Middle East, 249 -> Africa
PROFILE_RANGE = set(range(169, 221)) | set(range(249, 297))               # all template profile elements
COUNTRY_TABLES = {  # (heading-2, intro paragraph, table) per region in section 3.5
    "North America": (114, 115, 117), "Latin America": (119, 120, 122), "Europe": (124, 125, 127),
    "Asia Pacific": (130, 131, 133), "Middle East": (135, 136, 138), "Africa": (140, 141, 143),
}


# --------------------------------------------------------------------------- #
# XML helpers
# --------------------------------------------------------------------------- #
def _runs(p):
    return [r for r in p.findall(qn("w:r")) if r.find(qn("w:drawing")) is None and r.find(qn("w:fldChar")) is None
            and r.find(qn("w:instrText")) is None]


def _is_bold(r) -> bool:
    rpr = r.find(qn("w:rPr"))
    b = rpr.find(qn("w:b")) if rpr is not None else None
    return b is not None and b.get(qn("w:val")) not in ("0", "false")


def _text_run(proto, text: str, bold: bool | None = None):
    r = copy.deepcopy(proto)
    for ch in list(r):
        if ch.tag != qn("w:rPr"):
            r.remove(ch)
    if bold is not None:
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = r.makeelement(qn("w:rPr"), {})
            r.insert(0, rpr)
        for b in rpr.findall(qn("w:b")) + rpr.findall(qn("w:bCs")):
            rpr.remove(b)
        if bold:
            rpr.insert(0, rpr.makeelement(qn("w:b"), {}))
    t = r.makeelement(qn("w:t"), {})
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    return r


def set_text(p, text: str, bold: bool | None = None):
    """Replaces a paragraph's text, keeping its first plain run's formatting
    (or the first run's) and the paragraph properties (style, bullets)."""
    runs = _runs(p)
    plain = [r for r in runs if not _is_bold(r)]
    proto = (plain or runs or [p.makeelement(qn("w:r"), {})])[0]
    for r in runs:
        p.remove(r)
    p.append(_text_run(proto, text, bold))
    return p


def set_lead(p, lead: str, text: str):
    """Bold lead-in + plain text (template bullet/box pattern)."""
    runs = _runs(p)
    proto = (runs or [p.makeelement(qn("w:r"), {})])[0]
    for r in runs:
        p.remove(r)
    p.append(_text_run(proto, lead.rstrip() + " ", True))
    p.append(_text_run(proto, text, False))
    return p


def paras(el):
    return list(el.iter(qn("w:p")))


def _text_runs(p):
    return [r for r in _runs(p) if r.find(qn("w:t")) is not None]


def set_h1(p, num: str | None = None, title: str | None = None):
    """Chapter heading = number run ('03   ', orange) + title runs (white):
    each is rewritten in place so the two colours survive."""
    rs = _text_runs(p)
    if num is not None:
        rs[0].find(qn("w:t")).text = num
    if title is not None:
        for k, r in enumerate(rs[1:]):
            for t in r.findall(qn("w:t")):
                t.text = title if k == 0 else ""
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return p


def _style(el) -> str:
    s = el.find(".//" + qn("w:pStyle"))
    return s.get(qn("w:val")) if s is not None else ""


# Optional blocks of the template (body indices) and the section flag that keeps them.
def skipped(m: dict) -> tuple[set[int], list[str]]:
    sec = m["sections"]
    skip: set[int] = set()
    notes = []

    def drop(name: str, idx):
        skip.update(idx)
        notes.append(f"removed {name}: {sec[name.split(':')[0]]['why']}")

    profile_starts = [s for st in PROFILE_BLOCKS.values() for s in st]
    if not sec["type"]["show"]:
        drop("type: 'by type' sections", list(range(83, 88)) + list(range(158, 163)) + list(range(238, 243))
             + [s + o for s in profile_starts for o in range(2, 7)])
    if not sec["demand"]["show"]:
        drop("demand: 'by demand segment' sections", list(range(88, 94)) + list(range(163, 168))
             + list(range(243, 248)) + [s + o for s in profile_starts for o in range(7, 12)])
    if not sec["type"]["show"] and not sec["demand"]["show"]:
        skip.update({81, 82})
    if not sec["use_cases_me"]["show"]:
        drop("use_cases_me: Middle East use-case heat map", range(221, 225))
    if not sec["use_cases_af"]["show"]:
        drop("use_cases_af: Africa use-case heat map", range(298, 302))
    # vendor sections 7.2 / 7.3 are always kept: a donut, the positioning table, or one sentence saying
    # vendor-level data could not be confirmed (PLAN-05, vendor_shares)
    if not sec["capabilities"]["show"]:
        drop("capabilities: vendor capability map", range(343, 347))
    if not sec["dashboard"]["show"]:
        drop("dashboard: vendor dashboard", range(347, 351))
    if not sec["milestones"]["show"]:
        drop("milestones: regulatory milestones timeline", range(305, 308))
    if not sec["regulation"]["show"]:
        drop("regulation: Regulatory & Policy Outlook chapter", range(303, 325))
    return skip, notes


def to_verify_items(b: "Builder", ag: dict, m: dict) -> list[str]:
    """Every [to verify] item exactly as printed in the finished document (text
    and tables), plus the milestones marked on the timeline chart -- for the
    internal QA report. A sentence that only explains the marker is not an item."""
    explains = re.compile(r"items? (?:marked|flagged) \[to verify\]", re.I)
    items = []
    for el in b.out:
        if el.tag == qn("w:tbl"):
            for tr in el.iter(qn("w:tr")):
                cells = ["".join(t.text or "" for t in tc.iter(qn("w:t"))) for tc in tr.findall(qn("w:tc"))]
                if any("[to verify]" in c and not explains.search(c) for c in cells):
                    items.append(" | ".join(c for c in cells if c)[:200])
        else:
            text = "".join(t.text or "" for t in el.iter(qn("w:t")))
            items.extend(x[:200] for x in sentences(text) if "[to verify]" in x and not explains.search(x))
    if m["sections"]["regulation"]["show"] and m["sections"]["milestones"]["show"]:
        items += [f"Timeline chart: {x.get('date')} {x.get('jurisdiction', '')} -- {x.get('event')} [to verify]"
                  for x in (ag.get("regulation") or {}).get("milestones") or [] if x.get("verified") is False]
    return list(dict.fromkeys(items))


POSITIONING_HEADER = ["Vendor", "HQ", "Parent group", "Business type", "Role in this market"]
_POSITIONING_WIDTHS = [1900, 1200, 1700, 1600, 3106]      # the dashboard table's 9,506 twips


def _widen(tbl, widths: list[int]):
    """The template's 4-column vendor dashboard table with len(widths) columns
    (cells cloned from each row's last cell)."""
    grid = tbl.find(qn("w:tblGrid"))
    for gc in grid.findall(qn("w:gridCol")):
        grid.remove(gc)
    for w in widths:
        gc = grid.makeelement(qn("w:gridCol"), {})
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for tr in tbl.findall(qn("w:tr")):
        while len(tr.findall(qn("w:tc"))) < len(widths):
            last = tr.findall(qn("w:tc"))[-1]
            last.addnext(copy.deepcopy(last))
        for tc, w in zip(tr.findall(qn("w:tc")), widths):
            tcw = tc.find(qn("w:tcPr") + "/" + qn("w:tcW"))
            if tcw is not None:
                tcw.set(qn("w:w"), str(w))
    return tbl


def short_note(note: str | None) -> str:
    """A confirmed ownership note beside a vendor name ('now part of Huber'); a long
    sentence is not a name annotation and is left out."""
    note = str(note or "").strip().rstrip(".")
    return note if note and len(note.split()) <= 10 else ""


def positioning_rows(m: dict, vsafe) -> list[list[str]]:
    """vendor | HQ | parent group | business type | role -- confirmed vendors only, in alphabetical
    order; no tier or ranking (there are no sourced shares to rank by). The role is the one the product
    evidence line states. No MEA column: chapter 7 is the global view (changes.txt 2026-10-19)."""
    rows = []
    for r in m["vendor_positioning"]:
        own = short_note(r.get("ownership_note"))
        rows.append([r["vendor"] + (f" ({own})" if own and own.lower() not in r["vendor"].lower() else ""),
                     r["hq"], r["parent"], r["business_type"], vsafe(r["role"]).rstrip(".")])
    return rows


def tier_text(tiers: list[dict]) -> str:
    """The paragraph under the tier pyramid, from its data: who sits in each tier (the chart carries the shares)."""
    named = {t["name"]: ", ".join(t["players"]) for t in tiers if t.get("players")}
    if len(named) < 3:
        return ""
    return (f"Global value is led by a small group of Tier 1 manufacturers ({named['Tier 1']}), whose breadth of "
            f"portfolio and distribution reach set the pace of product launches. Tier 2 brings together established "
            f"multinationals and strong specialists ({named['Tier 2']}), while Tier 3 covers regional specialists "
            f"({named['Tier 3']}); the remaining value is spread across smaller regional and local manufacturers.")


def vendor_shares(b: "Builder", i: int, m: dict, s, root: str, charts: dict, vsafe) -> None:
    """Sections 7.2 / 7.3 (template elements 332-341), global view only (changes.txt 2026-10-19):
    the first slot is a global share donut where at least 3 shares are sourced; the second slot
    is the global positioning table (when vendors are confirmed). The MEA donut is not shown."""
    sec, vy = m["sections"], m["vendor_year"]
    show_g, pos = sec["vendor_global"]["show"], sec["positioning"]["show"]
    first = i < 337
    k = i - (332 if first else 337)                # 0 heading, 1 caption, 2 figure, 3 source, 4 paragraph
    if (m.get("vendor_unconfirmed") or {}).get("global"):
        if k == 0 and first:
            b.log.append("vendor section (Global) left out: too few confirmed vendors")
        return
    if first:
        if not show_g:
            return
        if k == 0:
            b.keep(i)
        elif k == 1:
            b.caption(i, "Figure", f"Vendor share of global {fmt.lower_name(root)} market value, {vy} (published figures)")
        elif k == 2:
            b.image(i, charts["vendor_global"])
        elif k == 3:
            pubs = dict.fromkeys(x.get("domain") or "" for x in m["vendor_sources"]["global"])
            b.para(i, "Source: " + ", ".join(p for p in pubs if p) + "; shares as stated by each source")
        else:
            b.para(i, s("global_share_para"))
        return
    if not pos:
        return
    if k == 0:
        b.para(i, "7.2 Vendor positioning (Global)")      # renumber() sets the real number
    elif k == 1:
        b.caption(i, "Table", f"Vendor positioning in the global {fmt.lower_name(root)} market")
    elif k == 2:
        b.src.append(_widen(b.el(349), _POSITIONING_WIDTHS))     # the dashboard table, 5 columns
        b.table(len(b.src) - 1, complete_rows(positioning_rows(m, vsafe), b.log, "vendor positioning"),
                header=POSITIONING_HEADER)
    elif k == 3:
        b.para(i, "Source: Company websites, filings and news; CMI analysis")
    elif not show_g:
        b.para(i, s("global_share_para"))


def internal_log(b: "Builder", m: dict, ag: dict, prose: dict, removed_notes: list[str],
                 verification: dict | None = None) -> dict:
    """What used to be appendix A.4-A.6 (the [to verify] list, the corrections made while
    building, the consistency log), for the internal QA report only. The counts agree: a
    sentence removed by the prose guards and one removed by the final sweep are counted apart."""
    tv = to_verify_items(b, ag, m)
    corr = list(m.get("corrections") or []) + list(removed_notes)
    dropped = prose.get("dropped") or {}
    n_drop = sum(len(v) for v in dropped.values())
    if n_drop:
        why = {"numbers not in model": 0, "names": 0, "repeat": 0, "headline restated": 0, "caveat": 0, "claim": 0}
        for v in dropped.values():
            for x in v:
                key = ("numbers not in model" if "numbers not in model" in x else "names" if "not in the data" in x
                       else "repeat" if x.startswith("[repeat]") else "headline restated" if "headline" in x
                       else "caveat" if "caveat" in x else "claim")
                why[key] += 1
        corr.append(f"{n_drop} sentence(s) removed by the prose guards while writing: " + ", ".join(
            f"{v} {k}" for k, v in {"with a number not in the master sheet": why["numbers not in model"],
                                    "with a ranking the register does not support": why["claim"],
                                    "naming a category absent from the data": why["names"],
                                    "repeating an earlier sentence": why["repeat"],
                                    "restating the headline outside the summary": why["headline restated"],
                                    "with template or draft wording": why["caveat"]}.items() if v))
    corr += [x for x in b.log if "incomplete row" in x or x.startswith((
        "empty heading removed", "figure left out", "table left out", "box left out", "vendor section",
        "regulatory snapshot"))]
    v = verification or {}
    log = ["Every value, share, gap and ranking was derived from the master data sheet, the number "
           "register (the CSV beside the report) and the rankings register before any text was written. " + fmt.PRECISION_RULE,
           "Checked: children sum to parents for every year (countries to sub-regions, sub-regions to regions, "
           "regions to global, segments to totals); every CAGR recomputed from its start and end values; every "
           "incremental value equals end minus start; every share uses its own denominator; every ranking and tie "
           "is taken from the rankings register; share sets sum to 100%."]
    if v:
        log.append(f"Number consistency sweep of the finished report: {v.get('checked', 0)} numbers checked; "
                   f"{v.get('first_sweep', 0)} text block(s) with a number not matching the register in the first "
                   f"sweep; {v.get('fixed', 0)} sentence(s) removed by the sweep (besides the {n_drop} removed by "
                   f"the prose guards); {v.get('final_sweep', 0)} mismatch(es) remaining after the final sweep.")
    log.append(f"{len(tv)} item(s) marked [to verify] in the document.")
    return {"to_verify": tv, "corrections": corr, "consistency_log": log,
            "prose_guard_removed": n_drop, "sweep_removed": v.get("fixed", 0)}


_NUMBERED = re.compile(r"^\s*\d+(?:\.\d+)*\s")


def heading_is_empty(level: int, text: str, nxt_level: int | None, nxt_text: str | None, at_end: bool) -> bool:
    """A heading has no body when the next block (blank paragraphs aside) is the end, or a heading of
    the same or a higher level. A numbered heading followed by an UNnumbered one of the same level
    ('3.5 Regional Analysis' -> 'North America ... by Country') is that heading's parent, not empty."""
    if at_end:
        return True
    if nxt_level is None or nxt_level > level:
        return False
    if nxt_level == level and _NUMBERED.match(text) and not _NUMBERED.match(nxt_text or ""):
        return False
    return True


def empty_headings(els: list) -> list[int]:
    """Indices (in `els`, the builder's body elements) of headings whose section has no body."""
    def level(el) -> int | None:
        mm = re.fullmatch(r"Heading(\d)", _style(el)) if el.tag == qn("w:p") else None
        return int(mm.group(1)) if mm else None

    def blank(el) -> bool:
        return el.tag == qn("w:p") and not el.xpath("string(.)").strip() and el.find(".//" + qn("w:drawing")) is None

    out = []
    for k, el in enumerate(els):
        lv = level(el)
        if lv is None:
            continue
        nxt = next((e for e in els[k + 1:] if not blank(e) and e.tag != qn("w:sectPr")), None)
        if heading_is_empty(lv, el.xpath("string(.)"), level(nxt) if nxt is not None else None,
                            nxt.xpath("string(.)") if nxt is not None else None, nxt is None):
            out.append(k)
    return out


def renumber(b: "Builder", toc_titles: dict[int, str], toc_pages: dict[int, int] | None) -> None:
    """Chapters and numbered sub-sections are renumbered in order (a removed
    section leaves no gap), and the contents page lists exactly the chapters
    present, with their new numbers."""
    new_of: dict[int, int] = {}
    ch = sub = sub3 = 0
    for el in b.out:
        st = _style(el)
        if st == "Heading1":
            rs = _text_runs(el)
            orig = int(re.match(r"\s*(\d+)", rs[0].find(qn("w:t")).text).group(1))
            ch, sub = ch + 1, 0
            new_of[orig] = ch
            set_h1(el, num=f"{ch:02d}   ")
        elif st == "Heading2":
            text = el.xpath("string(.)")
            mm = re.match(r"\s*\d+\.\d+\s+(.*)", text)
            if mm and ch:
                sub, sub3 = sub + 1, 0
                set_text(el, f"{ch}.{sub} {mm.group(1)}")
        elif st == "Heading3":
            # 2.1.1 Global / 2.1.2 Middle East & Africa (changes.txt 2026-10-19)
            mm = re.match(r"\s*\d+\.\d+\.\d+\s+(.*)", el.xpath("string(.)"))
            if mm and ch and sub:
                sub3 += 1
                set_text(el, f"{ch}.{sub}.{sub3} {mm.group(1)}")
    for orig, el in b.toc.items():
        if orig not in new_of:
            b.out.remove(el)
            continue
        n = new_of[orig]
        rs = _text_runs(el)
        rs[0].find(qn("w:t")).text = f" {n}    "
        for k, r in enumerate(rs[1:-1]):
            for t in r.findall(qn("w:t")):
                t.text = toc_titles[orig] if k == 0 else ""
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        if toc_pages and n in toc_pages:
            rs[-1].find(qn("w:t")).text = str(toc_pages[n])


# --------------------------------------------------------------------------- #
# builder
# --------------------------------------------------------------------------- #
class Builder:
    def __init__(self, template: Path):
        self.doc = docx.Document(str(template))
        self.body = self.doc.element.body
        self.src = [copy.deepcopy(e) for e in self.body.iterchildren()]
        self.out: list = []
        self.fig = self.tbl = 0
        self.pic_id = 9000
        self.log: list[str] = []
        self.toc: dict[int, object] = {}      # original chapter number -> its contents-page entry

    def el(self, i: int):
        return copy.deepcopy(self.src[i])

    def _source_of_skipped_figure(self, text: str) -> bool:
        """True for the 'Source: ...' line right after a figure or table that was left out, or one
        that would sit directly under a heading (nothing above it to cite)."""
        skip, self.skip_source = getattr(self, "skip_source", False), False
        if not text.lstrip().startswith("Source"):
            return False
        last = self.out[-1] if self.out else None
        under_heading = last is not None and last.tag == qn("w:p") and _style(last).startswith("Heading")
        return skip or under_heading

    def _drop_caption(self, kind: str, what: str):
        """A figure or table with nothing to show: its caption (and number) goes, and the source
        line that follows is skipped; logged for the internal QA report."""
        last = self.out[-1] if self.out else None
        cap = last.xpath("string(.)").strip() if last is not None and last.tag == qn("w:p") else ""
        if re.match(rf"{kind} \d+:", cap):
            self.out.pop()
            if kind == "Figure":
                self.fig -= 1
            else:
                self.tbl -= 1
        self.skip_source = True
        why = "no data to chart" if kind == "Figure" else "no data"
        self.log.append(f"{kind.lower()} left out ({why}): {re.sub(rf'^{kind} \d+:\s*', '', cap) or what}")

    def keep(self, i: int):
        el = self.el(i)
        if self._source_of_skipped_figure(el.xpath("string(.)")):
            return
        self.out.append(el)

    def para(self, i: int, text: str | None, bold: bool | None = None):
        if isinstance(text, list):              # a paragraph slot the model returned as a list of sentences
            text = " ".join(str(x).strip() for x in text if isinstance(x, str) and x.strip())
        if self._source_of_skipped_figure(text or ""):
            return
        if text:
            self.out.append(set_text(self.el(i), text, bold))
        else:
            self.log.append(f"element {i}: no text -- left out")

    def heading3(self, title: str):
        """A numbered sub-section ('0.0.1 Global'; renumber() sets the real number), cloned from the
        template's Heading3 (element 83)."""
        self.out.append(set_text(self.el(83), f"0.0.1 {title}"))

    def lead(self, i: int, lead: str, text: str):
        self.out.append(set_lead(self.el(i), lead, text))

    def bullets(self, i: int, items: list[str]):
        for it in items:
            self.out.append(set_text(self.el(i), it, False))

    def caption(self, i: int, kind: str, text: str):
        if kind == "Figure":
            self.fig += 1
            n = self.fig
        else:
            self.tbl += 1
            n = self.tbl
        self.out.append(set_text(self.el(i), f"{kind} {n}: {text}", True))

    def image(self, i: int, path: Path | None):
        """Places a chart. `path` None = the chart was skipped (charts.py: nothing to draw, e.g. a
        zero-value unit): the figure is left out cleanly -- its caption (and figure number) and its
        source line go too, and the omission is logged."""
        if path is None:
            self._drop_caption("Figure", f"element {i}")
            return
        el = self.el(i)
        r_id, _ = self.doc.part.get_or_add_image(str(path))
        with Image.open(path) as im:
            w, h = im.size
        for blip in el.iter(qn("a:blip")):
            blip.set(qn("r:embed"), r_id)
        for tag in (qn("wp:extent"), qn("a:ext")):
            for ext in el.iter(tag):
                cx = int(ext.get("cx"))
                cy = int(cx * h / w)
                if cy > 5_200_000:                      # very tall: shrink to fit the page
                    cx, cy = int(5_200_000 * w / h), 5_200_000
                ext.set("cx", str(cx))
                ext.set("cy", str(cy))
        # alt text = this figure's caption (the template's says 'drone'; Word copies it into the PDF)
        last = self.out[-1] if self.out else None
        cap = last.xpath("string(.)").strip() if last is not None and last.tag == qn("w:p") else ""
        alt = re.sub(r"^Figure \d+:\s*", "", cap) if cap.startswith("Figure") else path.stem.replace("_", " ")
        for tag in (qn("wp:docPr"), qn("pic:cNvPr")):
            for node in el.iter(tag):
                self.pic_id += 1
                node.set("id", str(self.pic_id))
                node.set("name", path.stem)
                node.set("descr", alt)
                node.set("title", alt)
        self.out.append(el)

    def table(self, i: int, rows: list[list[str]], total: list[str] | None = None, header: list[str] | None = None):
        """Header row kept; data rows cloned from the template's first data
        row (and its last row for a bold total). A cell whose value the
        template also uses (High / Medium / Low) takes that template cell's
        shading, so rating colours follow the value, not the row position. No rows = no data:
        the table, its caption and its source line are left out (like image(None))."""
        if not rows:
            self._drop_caption("Table", f"table at element {i}")
            return
        tbl = self.el(i)
        trs = tbl.findall(qn("w:tr"))
        proto, last = trs[1], trs[-1]
        shading, para_proto = {}, {}
        for tr in trs[1:]:
            for tc in tr.findall(qn("w:tc")):
                pr = tc.find(qn("w:tcPr"))
                shd = pr.find(qn("w:shd")) if pr is not None else None
                key = tc.xpath("string(.)").strip()
                shading.setdefault(key, shd)
                para_proto.setdefault(key, tc.find(qn("w:p")))
        for tr in trs[1:]:
            tbl.remove(tr)
        if header:
            for tc, val in zip(trs[0].findall(qn("w:tc")), header):
                set_text(tc.find(qn("w:p")), val, True)
        for row in rows + ([total] if total else []):
            tr = copy.deepcopy(last if (total and row is total) else proto)
            for tc, val in zip(tr.findall(qn("w:tc")), row):
                ps = tc.findall(qn("w:p"))
                set_text(ps[0], str(val), True if (total and row is total) else None)
                for extra in ps[1:]:
                    tc.remove(extra)
                if str(val) in shading and str(val) in ("High", "Medium", "Low"):
                    # the template's own cell for this rating: its text colour goes with its shading
                    tc.replace(tc.find(qn("w:p")), copy.deepcopy(para_proto[str(val)]))
                    pr = tc.find(qn("w:tcPr"))
                    if pr is not None:
                        for old in pr.findall(qn("w:shd")):
                            pr.remove(old)
                        if shading[str(val)] is not None:
                            pr.append(copy.deepcopy(shading[str(val)]))
            tbl.append(tr)
        self.out.append(tbl)

    def box(self, i: int, title: str | None, items: list, numbered: bool = False):
        """Single-cell callout box: title paragraph + item paragraphs. No items = no box."""
        items = [x for x in items or [] if x]
        if not items:
            self.log.append(f"box left out (no content): {title or f'element {i}'}")
            return
        tbl = self.el(i)
        ps = paras(tbl)
        cell = ps[0].getparent()
        if title:
            set_text(ps[0], title, True)
        proto = ps[1]
        for p in ps[1:]:
            cell.remove(p)
        for k, it in enumerate(items, start=1):
            p = copy.deepcopy(proto)
            if isinstance(it, tuple):
                set_lead(p, it[0], it[1])
            elif numbered:
                # the box numbers its items; drop a number the text already carries ("1. 1. ...")
                set_lead(p, f"{k}.", re.sub(r"^\s*(?:\d+[.)]\s*)+", "", str(it)))
            else:
                set_text(p, it, False)
            cell.append(p)
        self.out.append(tbl)

    def drop_empty_headings(self):
        """Removes every heading whose section ended up with no body (repeated, so a chapter left
        with only empty sub-sections goes too); logged for the internal QA report."""
        while True:
            idx = empty_headings(self.out)
            if not idx:
                return
            for k in reversed(idx):
                self.log.append(f"empty heading removed: {self.out[k].xpath('string(.)').strip()[:80]}")
                del self.out[k]

    def save(self, path: Path, header: str, footer_label: str):
        for c in list(self.body.iterchildren()):
            self.body.remove(c)
        for e in self.out:
            self.body.append(e)
        # 'Middle East & Africa (MEA)' everywhere in the text, contents and headings (user decision 2026-10-20);
        # a name split across two runs is joined into the first run so it is caught too
        for p in self.body.iter(qn("w:p")):
            ts = [t for t in p.iter(qn("w:t"))]
            if "Africa" not in "".join(t.text or "" for t in ts):
                continue
            for a, nxt in zip(ts, ts[1:]):
                joined = (a.text or "") + (nxt.text or "")
                if re.search(r"Middle East\s*(?:&|and)\s*$", a.text or "") and (nxt.text or "").startswith("Africa"):
                    a.text, nxt.text = joined[:len(joined) - len(nxt.text) + len("Africa")], nxt.text[len("Africa"):]
            for t in ts:
                if t.text and "Africa" in t.text:
                    t.text = fmt.mea(t.text)
        for s in self.doc.sections:
            for p in s.header.paragraphs:
                for r in p.runs:
                    if "Outlook" in (r.text or ""):
                        r.text = header
            for p in s.footer.paragraphs:
                for r in p.runs:
                    if re.search(r"confidential|template|whitepaper", r.text or "", re.I):
                        r.text = footer_label
        # the document properties are this market's, not the template's: Word's PDF export uses the
        # title as the PDF title, which viewers show instead of the file name ('The Drone Market Outlook')
        cp = self.doc.core_properties
        cp.title = header
        cp.subject = header
        cp.comments = header
        cp.keywords = ""
        cp.author = cp.last_modified_by = "Coherent Market Insights"
        cp.created = cp.modified = datetime.now()
        xml = self.doc.element.xml
        for r_id, rel in list(self.doc.part.rels.items()):
            if "image" in rel.reltype and f'"{r_id}"' not in xml:
                del self.doc.part.rels[r_id]
        app = next((p for p in self.doc.part.package.parts if str(p.partname) == "/docProps/app.xml"), None)
        if app is not None:          # 'TitlesOfParts' still names the template's title
            app._blob = re.sub(rb"(<vt:lpstr>)[^<]*Outlook[^<]*(</vt:lpstr>)",
                               lambda mm: mm.group(1) + header.replace("&", "&amp;").encode() + mm.group(2), app.blob)
        self.doc.save(str(path))


# --------------------------------------------------------------------------- #
# content
# --------------------------------------------------------------------------- #
def _root(market: str) -> str:
    return re.sub(r"^\s*global\s+|\s+market\s*$", "", market, flags=re.I).strip()


def build(template: Path, out_path: Path, m: dict, prose: dict, ag: dict, charts: dict[str, Path],
          definition: dict | None, toc_pages: dict[int, int] | None = None,
          verification: dict | None = None, internal: dict | None = None) -> list[str]:
    """Builds the client .docx. `internal`, when given, is filled with the internal QA log
    (to-verify items, corrections, consistency log) that the client document does not carry."""
    b = Builder(template)
    S = prose.get("slots", {})
    root = _root(m["market"])
    tdim, ddim = m["dimensions"]["type"], m["dimensions"]["demand"]
    tl, dl = tdim["singular"].lower(), ddim["singular"].lower()
    g, mea = m["global"], m["regions"]["Middle East & Africa"]
    me, af = m["subregions"]["Middle East"], m["subregions"]["Africa"]
    ex = ag.get("_extras") or {}                  # extras.derive (changes.txt 2026-10-19)
    sources = [(l, it) for l, it in ag.get("_sources") or [] if it]         # resources.build (9.1 Sources)
    greg = ex.get("global_regulation") or {}      # 6.1 Global (extras.global_regulation)

    from AiAgentsFleet.cmi_report.register import FactGuard, build_register
    register = build_register(m, ag)
    ng = FactGuard(register)                      # agent-written cells obey the same register
    from AiAgentsFleet.cmi_report.prose import make_guard, set_template
    set_template(template)
    cg = make_guard(m, None)
    dropped: list[str] = []
    absent = absent_names(m)

    def safe(text: str | None) -> str:
        """Agent-written table/box text through the same guards as the prose."""
        return clean_text(text or "", ng, cg, dropped, absent) if text else ""

    def vsafe(text: str | None) -> str:
        """Vendor-chapter text: the same guards, then the vendor data (guard.vendor_clean)."""
        return vendor_clean(safe(text), m.get("vendor_text"), dropped) if text else ""

    def s(key: str, fallback: str | None = None):
        v = S.get(key)
        return v if v else fallback

    def lst(key: str) -> list:
        v = S.get(key)
        return v if isinstance(v, list) else []

    def obj(key: str) -> dict:
        v = S.get(key)
        return v if isinstance(v, dict) else {}

    def pairs(key: str) -> list[tuple[str, str]]:
        """(lead, text) items from a slot, whatever shape the model returned:
        {lead, text} objects, 'Lead: text' strings; anything else is skipped."""
        out = []
        for x in lst(key):
            if isinstance(x, dict) and x.get("lead") and x.get("text"):
                out.append((str(x["lead"]).strip(), str(x["text"]).strip()))
            elif isinstance(x, str) and ":" in x[:60]:
                lead, text = x.split(":", 1)
                out.append((lead.strip(), text.strip()))
            else:
                b.log.append(f"{key}: an item without lead/text was left out")
        return [(l, t) for l, t in out if l and t]

    seg = m["segments"]
    sec = m["sections"]
    skip, removed_notes = skipped(m)
    # the scorecard's "fastest-growing" row uses a dimension the report actually shows
    fdim = "type" if sec["type"]["show"] else "demand" if sec["demand"]["show"] else None
    fastest = {geo: max(seg[geo][fdim], key=lambda n: seg[geo][fdim][n]["cagr"]) for geo in seg} if fdim else {}
    fdim_label = (tdim if fdim == "type" else ddim)["singular"].lower() if fdim else ""
    toc_titles: dict[int, str] = {}
    size_unsourced = (m.get("size_check") or {}).get("status") != "sourced"
    from AiAgentsFleet.cmi_report.model import regulatory_countries
    model_reg_countries = regulatory_countries(ag)
    # regulators and requirements as fact-checked (verify_facts.apply); a row whose authority or
    # requirements were not confirmed carries no information and is left out
    _unconf = re.compile(r"\bnot confirmed\b|\bunconfirmed\b|\bnot (?:identified|available)\b", re.I)
    reg_rows = complete_rows(
        [[r.get("country", ""), r.get("regulator", ""), safe(r.get("requirements"))]
         for r in (ag.get("regulation") or {}).get("snapshot") or []
         if canonical(r.get("country", "")) in model_reg_countries], b.log, "regulatory snapshot")
    n_reg = len(reg_rows)
    reg_rows = [r for r in reg_rows if not any(_unconf.search(str(c)) for c in r[1:])]
    if len(reg_rows) < n_reg:
        b.log.append(f"regulatory snapshot: {n_reg - len(reg_rows)} row(s) without confirmed authority or "
                     f"requirements left out")

    def seg_para(geo: str, kind: str) -> str | None:
        # the written analysis, else a one-to-two-line insight from the ranking (no figures repeated)
        dim = (tdim if kind == "type" else ddim)["singular"]
        return s(f"{kind}_para__{geo}") or seg_insight(seg[geo]["type" if kind == "type" else "demand"], dim, geo)

    def country_rows(region: str) -> tuple[list[list[str]], list[str]]:
        d = m["countries"][region]
        rest = [k for k in d if k.startswith("Rest of")]
        order = sorted((k for k in d if k not in rest), key=lambda k: -d[k]["v26"]) + rest
        unit = fmt.unit_for([d[k]["v26"] for k in d])
        rows = [[k, fmt.usd_in(d[k]["v26"], unit), fmt.usd_in(d[k]["v33"], unit), fmt.pct(d[k]["cagr"]),
                 fmt.incr(d[k]["incr"], unit)] for k in order]
        tot = m["subregions"].get(region) or m["regions"].get(region)
        total = [f"{region} total", fmt.usd_in(tot["v26"], unit), fmt.usd_in(tot["v33"], unit),
                 fmt.pct(tot["cagr"]), fmt.incr(tot["incr"], unit)]
        return rows, total

    def region_mix(b, dim: str, label: str):
        """Comparative chart by region for one segment dimension (changes.txt 2026-10-19)."""
        if not charts.get(f"mix_{dim}"):
            return
        b.caption(84, "Figure", f"{label} mix by region, {Y0}E")
        b.image(85, charts[f"mix_{dim}"])
        b.para(86, "Source: CMI analysis")

    regions_ch = regional_chapters(m)

    def regional_chapter(r: str, orig: int):
        """One region's opportunity landscape, with every chart of the global chapter for that region
        (changes.txt 2026-10-19): countries over the years, scenarios, segments and their country mix,
        country shares, absolute opportunity, the country table and the country profiles."""
        d, name = m["regions"][r], fmt.lower_name(root)
        b.out.append(set_h1(b.el(146), num=f"{orig}   ", title=f"{r} {root} Market: Opportunity Landscape"))
        lead_in = (f"The {r} {name} market is estimated at {fmt.usd(d['v26'])} in {Y0}E and is forecast to reach "
                   f"{fmt.usd(d['v33'])} by {Y1}F ({fmt.pct(d['cagr'])} CAGR).")
        b.para(147, open_chapter(s(f"region_para__{r}"), rf"\b{re.escape(r)}\b", lead_in))
        b.caption(148, "Figure", f"{r} {name} market by country, {Y0 - 1}A-{Y1}F "
                                 f"(USD {'billion' if d['v33'] >= 1000 else 'million'})")
        b.image(149, charts.get(f"stack_{r}"))
        b.para(150, "Source: CMI analysis")
        b.para(69, f"0.1 {r} scenario analysis")
        b.caption(70, "Figure", f"{r} {name} market scenarios: conservative, base and optimistic")
        b.image(71, charts.get(f"scen_{r}"))
        b.para(72, "Source: CMI analysis")
        if r in seg and (charts.get(f"type_{r}") or charts.get(f"demand_{r}")):
            b.para(81, f"0.2 {r} {root} Market Size by Segments")
            for kind, dim in (("type", tdim), ("demand", ddim)):
                key = f"{kind}_{r}"
                if not sec[kind]["show"] or not charts.get(key):
                    continue
                b.para(83, f"{r} Opportunity by {dim['singular']}")
                b.caption(84, "Figure", f"{r} market by {dim['singular'].lower()}, {Y0}E vs {Y1}F"
                                        + (" (USD)" if kind == "demand" else ""))
                b.image(85, charts[key])
                b.para(86, "Source: CMI analysis")
                b.para(87, seg_para(r, kind))
                if charts.get(f"cmix_{kind}_{r}"):
                    b.caption(84, "Figure", f"{dim['singular']} mix by country, {Y0}E")
                    b.image(85, charts[f"cmix_{kind}_{r}"])
                    b.para(86, "Source: CMI analysis")
        b.para(94, f"0.3 {r} {root} Market Size by Country")
        b.caption(95, "Figure", f"Country share of the {r} market, {Y0 - 1}A")
        b.image(98, charts.get(f"cshare_{r}"))
        b.para(99, "Source: CMI analysis")
        b.caption(103, "Figure", f"{r}: Absolute dollar opportunity by country, {Y0}-{Y1}")
        b.image(104, charts.get(f"copp_{r}"))
        b.para(105, "Source: CMI analysis")
        rows, total = country_rows(r)
        b.caption(152, "Table", f"{r} {name} market by country (USD)")
        b.table(153, rows, total, header=COUNTRY_HEADER)
        b.para(154, "Source: CMI analysis")
        profiles, why = obj("country_profiles"), m.get("geography_why") or {}
        for c, v in m["countries"][r].items():
            text = with_lead(c, profiles.get(c) or safe(why.get(c)), v, d, r, m["countries"][r])
            b.lead(29, f"{c}:", text.strip())

    month = datetime.now().strftime("%B %Y").upper()
    practice = str(s("practice") or "Industry Research Practice").strip().rstrip(".")
    # no chapter-closing bridge sentences ('The next chapter turns to ...'): the template has none (2026-10-16)
    for i in range(len(b.src)):
        if i in REMOVE or i in skip:
            continue
        # 3.5 'Regional Analysis, By Key Countries' is left out: every region's country table and profiles
        # are in its own chapter now; 7.1 'Value chain analysis' is removed (user decisions 2026-10-19)
        if 112 <= i <= 145 or 327 <= i <= 331:
            continue
        if i == 146:
            for k, r in enumerate(regions_ch):
                regional_chapter(r, REGION_CH + k)
        # ------------------------------------------------------------ cover / TOC / boxes
        if i == 0:
            b.image(i, charts["cover"])
        elif i == 1:
            el = b.el(i)
            rs = _runs(el)
            for r in rs:
                el.remove(r)
            el.append(_text_run(rs[0], month, True))      # no "WHITEPAPER" label on the cover (2026-10-20)
            b.out.append(el)
        elif i == 2:
            b.para(i, f"{root} Market Outlook {Y0}-{Y1}")
        elif i == 6:
            b.para(i, practice)
        elif 9 <= i <= 18:
            # contents entry for original chapter i-8; numbers, titles and pages are
            # filled by renumber() once we know which chapters this report has
            titles = {9: "Executive Summary: Global vs Middle East & Africa",
                      10: "Current Market Trends, Growth Factors, Restraints and Impact",
                      11: f"Global {root} Market Size: Scenario Analysis by Segments, Region & Top Countries",
                      12: f"Middle East {root} Market Size: Opportunity Analysis by Segments & Country",
                      13: f"Africa {root} Market Size: Opportunity Analysis by Segments & Country",
                      14: "Regulatory & Policy Outlook: Global vs Middle East & Africa",
                      15: "Vendor Intelligence: Global"}
            if i == 9:
                # the report starts with the market definition, before the summary (user decision 2026-10-20)
                el = b.el(i)
                toc_titles[DEFINITION_CH] = DEFINITION_TITLE
                b.toc[DEFINITION_CH] = el
                b.out.append(el)
            if i == 12:
                # contents entries for the regional chapters, before the Middle East (original numbers 20+)
                for k, r in enumerate(regions_ch):
                    el = b.el(i)
                    toc_titles[REGION_CH + k] = f"{r} {root} Market Size: Opportunity Analysis by Segments & Country"
                    b.toc[REGION_CH + k] = el
                    b.out.append(el)
            el = b.el(i)
            texts = _text_runs(el)
            toc_titles[i - 8] = titles.get(i) or "".join(r.xpath("string(.)") for r in texts[1:-1])
            b.toc[i - 8] = el
            b.out.append(el)
        elif i == 20:
            # A1: the market definition comes first and everything else obeys it; inclusions and
            # exclusions are itemised once, in the Methodology scope table
            dfn = m["definition"]
            what = _sentence(dfn["definition"]) if dfn.get("definition") else ""
            # what is counted and at which value-chain level, verbatim from the definition (PLAN-04)
            measures = _sentence(dfn["measures"]) if dfn.get("measures") and dfn["measures"] != dfn.get("definition") \
                else ""
            level = f"Value-chain level: {dfn['value_chain_level']}." if dfn.get("value_chain_level") else ""
            units = (f"USD at current prices; {Y0 - 1}A is the base year (actual), {Y0}E the estimate year and "
                     f"{Y0 + 1}F-{Y1}F the forecast period (A = Actual, E = Estimated, F = Forecast).")
            # chapter 01 says what the market IS, what it covers, includes and excludes (CEO feedback 2026-10-20)
            b.out.append(set_h1(b.el(23), num=f"{DEFINITION_CH:02d}   ", title=DEFINITION_TITLE))
            b.para(24, safe(s("definition_para")) or what)
            # no 'Market coverage' list: the taxonomy figure covers segments and geography (CEO feedback
            # 2026-10-20); the definition, inclusions, exclusions and taxonomy remain

            def items(key: str, text_key: str) -> list[str]:
                xs = dfn.get(key) or [x for x in re.split(r";\s*", dfn.get(text_key) or "") if x.strip()]
                return [_sentence(x.strip()[0].upper() + x.strip()[1:]) for x in xs if x.strip()]
            if items("included_items", "included"):
                b.para(49, "0.2 Inclusions")
                b.bullets(29, items("included_items", "included"))
            if items("excluded_items", "excluded"):
                b.para(49, "0.3 Exclusions")
                b.bullets(29, items("excluded_items", "excluded"))
            b.para(49, "0.4 Market taxonomy")
            # the market taxonomy follows the definition (changes.txt 2026-10-19)
            b.caption(25, "Figure", f"{root} market taxonomy")
            b.image(26, charts.get("taxonomy"))
            b.para(48, "Source: CMI analysis")
        # ------------------------------------------------------------ chapter 1
        elif i == 24:
            headline = (f"CMI estimates the global {fmt.lower_name(root)} market at {fmt.usd(g['v25'])} in "
                        f"{Y0 - 1}, rising to {fmt.usd(g['v26'])} in {Y0} and {fmt.usd(g['v33'])} by "
                        f"{Y1}, a {fmt.pct(g['cagr'])} CAGR.")
            opener = s("exec_para", headline)
            # an unsourced headline is always called a CMI estimate (PLAN-02)
            if size_unsourced and not re.search(r"\bCMI estimat", opener):
                opener = headline
            # the opening paragraph is about the global market only; MEA has its own snapshot in 1.2
            # (CEO feedback 2026-10-20)
            b.para(i, global_only(((s("thesis") or "") + " " if s("thesis") else "") + opener) or headline)
        elif i == 25:
            b.caption(i, "Figure", "Headline market metrics: Global vs Middle East & Africa")
        elif i == 26:
            b.image(i, charts["kpi"])
        elif i in (29, 33):
            b.bullets(i, lst("global_bullets" if i == 29 else "mea_bullets"))
        elif i in (30, 31, 34, 35):
            continue
        elif i == 36:
            from AiAgentsFleet.cmi_report import rules
            pace = ("in line with" if abs(mea["cagr"] - g["cagr"]) < rules.TIE_PP else
                    "faster than" if mea["cagr"] > g["cagr"] else "more slowly than")
            b.caption(i, "Figure", f"MEA grows {pace} the global market ({Y0}-{Y1} CAGR); MEA share of the global "
                                   f"market, {Y0 - 1}A")
        elif i == 37:
            b.image(i, charts["cagr_share"])
        elif i == 40:
            b.caption(i, "Table", "Global vs MEA scorecard")
        elif i == 41:
            dyn = ag.get("dynamics") or {}
            fleet_dr = m.get("fleet_drivers") or {}

            def top(lst_, key):
                return max(lst_, key=lambda x: abs(float(x.get(key) or 0)))["name"] if lst_ else None

            def largest_pool(geo):
                d = seg[geo]["demand"] if sec["demand"]["show"] and seg[geo]["demand"] else \
                    seg[geo]["type"] if sec["type"]["show"] else {}
                if not d:
                    return None
                n = max(d, key=lambda k: d[k]["v26"])
                return f"{n} ({fmt.usd(d[n]['v26'])} in {Y0}E)"

            # data-derived values: a scorecard cell is never blank
            fallback = {
                "largest_demand_pool": (largest_pool("Global"), largest_pool("Middle East") and
                                        f"Middle East: {largest_pool('Middle East')}; Africa: {largest_pool('Africa')}"),
                "primary_catalyst": (top(fleet_dr.get("drivers") or [], "cagr_contribution_bps"),
                                     top(dyn.get("mea_drivers") or [], "impact_pp")),
                "primary_constraint": (top(fleet_dr.get("restraints") or [], "cagr_drag_bps"),
                                       top(dyn.get("mea_restraints") or [], "impact_pp")),
            }

            def two(key):
                # a row with no data for either column is left out (no 'Not applicable' cells)
                o = obj(key)
                fg, fm = fallback.get(key, (None, None))
                return [safe(o.get("global")) or fg or "", safe(o.get("mea")) or fm or ""]
            rows = [[f"Market size {Y0 - 1}A", fmt.usd(g["v25"]), fmt.usd(mea["v25"])],
                    [f"Market size {Y0}E", fmt.usd(g["v26"]), fmt.usd(mea["v26"])],
                    [f"Market size {Y1}F", fmt.usd(g["v33"]), fmt.usd(mea["v33"])],
                    [f"CAGR {Y0}-{Y1}", fmt.pct(g["cagr"]), fmt.pct(mea["cagr"])],
                    [f"Incremental revenue {Y0}-{str(Y1)[2:]}", fmt.usd(g["incr"]), fmt.usd(mea["incr"])],
                    ["Largest demand pool", *two("largest_demand_pool")]]
            if fdim:
                ft = seg["Global"][fdim][fastest["Global"]]
                fme, faf = fastest["Middle East"], fastest["Africa"]
                cme, caf = seg["Middle East"][fdim][fme]["cagr"], seg["Africa"][fdim][faf]["cagr"]
                mea_fast = (f"{fme} (Middle East {fmt.pct(cme)}; Africa {fmt.pct(caf)})" if fme == faf else
                            f"{fme} (Middle East {fmt.pct(cme)}); {faf} (Africa {fmt.pct(caf)})")
                from AiAgentsFleet.cmi_report import rules

                def lead(geo):          # near-ties (rules.TIE_PP) are named together, never one "fastest"
                    tied = rules.top_tied(seg[geo][fdim], "cagr")
                    return " and ".join(tied) + " (similar)" if len(tied) > 1 else tied[0]
                if any(len(rules.top_tied(seg[geo][fdim], "cagr")) > 1 for geo in ("Global", "Middle East", "Africa")):
                    mea_fast = (f"{lead('Middle East')} (Middle East {fmt.pct(cme)}); "
                                f"{lead('Africa')} (Africa {fmt.pct(caf)})")
                rows.append([f"Fastest-growing {fdim_label}", f"{lead('Global')} ({fmt.pct(ft['cagr'])})", mea_fast])
            rows += [["Primary growth catalyst", *two("primary_catalyst")],
                     ["Primary constraint", *two("primary_constraint")]]
            b.table(i, complete_rows(rows, b.log, "scorecard"))
        elif i == 43:
            # at least three conclusions (the prose's, topped up from the body findings), or no box
            concl = fill_conclusions(lst("key_conclusions"), S)
            if not concl:
                b.log.append("Key conclusions: fewer than three findings -- box left out")
                continue
            b.box(i, "Key conclusions", concl, numbered=True)
        # ------------------------------------------------------------ chapter 2
        elif i == 45:
            b.para(i, open_chapter(s("trends_intro"), r"(?i)trend|force|shift",
                                   f"Six forces are reshaping the {fmt.lower_name(root)} market, and each one shows up in the "
                                   f"regional and segment numbers that follow."))
        elif i == 46:
            b.caption(i, "Figure", f"Six trends shaping the {fmt.lower_name(root)} market")
        elif i == 47:
            b.image(i, charts["trends"])
        elif i == 50:
            # every trend shown in the figure is explained in the text: 2.1.1 the global view,
            # 2.1.2 how the trends play out in MEA (changes.txt 2026-10-19)
            b.heading3("Global")
            got = {l.lower().rstrip("."): t for l, t in pairs("trend_paras")}
            for tr in ((ag.get("dynamics") or {}).get("trends") or [])[:6]:
                title = tr.get("title", "")
                text = got.get(title.lower()) or next((t for l, t in got.items() if title.lower()[:12] in l), None) \
                    or _sentence(tr.get("tagline", ""))
                b.lead(i, title.rstrip(".") + ".", text)
            mea_trends = pairs("trend_paras_mea")
            if mea_trends:
                b.heading3("Middle East & Africa")
                for lead, text in mea_trends:
                    b.lead(i, lead.rstrip(".:") + ".", text)
        elif i in (51, 52):
            continue
        elif i == 54:
            # 2.2.1 global factors (fleet desk research, contributions to the global CAGR), 2.2.2 MEA
            if charts.get("drivers_global"):
                b.heading3("Global")
                b.para(i, s("drivers_para_global"))
                b.caption(56, "Figure", "Contribution of growth factors and restraints to the global CAGR "
                                        "(percentage points)")
                b.image(57, charts["drivers_global"])
                b.para(58, "Source: CMI analysis")
                b.heading3("Middle East & Africa")
            b.para(i, s("drivers_para"))
        elif i == 56:
            basis = (ag.get("dynamics") or {}).get("impact_basis")
            b.caption(i, "Figure", "Relative impact of growth factors and restraints in MEA (score 1-5)"
                      if basis == "score" else
                      "Impact of growth factors and restraints on MEA CAGR vs the global CAGR (percentage points)"
                      if basis == "pp_vs_global" else
                      "Impact of growth factors and restraints on MEA CAGR (percentage points)")
        elif i == 58:
            # the template's 'expert-panel scoring' source line describes work that never happened (X16)
            b.para(i, "Source: CMI analysis")
        elif i == 57:
            b.image(i, charts["drivers"])
        elif i == 60:
            # 2.3.1 global matrix (fleet factors), 2.3.2 the MEA matrix
            if ex.get("global_matrix"):
                b.heading3("Global")
                b.para(i, s("matrix_para_global"))
                b.caption(61, "Table", "Impact assessment of global market factors")
                b.table(62, complete_rows(ex["global_matrix"], b.log, "global impact matrix"))
                b.para(63, "Source: CMI analysis")
                b.heading3("Middle East & Africa")
            b.para(i, s("matrix_para"))
        elif i == 61:
            b.caption(i, "Table", "Impact assessment of market factors in the Middle East & Africa")
        elif i == 62:
            from AiAgentsFleet.cmi_report import rules
            dyn = ag.get("dynamics") or {}
            # the impact label follows the reconciled impacts the report prints (no 'Low' factor with a top pp)
            rows = [[safe(r.get("factor")), r.get("type", ""),
                     rules.impact_label(str(r.get("factor") or ""), dyn) or r.get("impact", ""), r.get("timeframe", ""),
                     safe(r.get("most_exposed"))] for r in dyn.get("impact_matrix") or []]
            b.table(i, rows)
        # ------------------------------------------------------------ chapter 3
        elif i == 64:
            b.out.append(set_h1(b.el(i), title=f"Global {root} Market Size: Scenario Analysis"))
        elif i == 65:
            b.para(i, s("sizing_para", f"CMI's base case values the global {fmt.lower_name(root)} market at "
                                       f"{fmt.usd(g['v26'])} in {Y0}E, reaching {fmt.usd(g['v33'])} in {Y1}F."))
        elif i == 66:
            b.caption(i, "Figure", f"Global {fmt.lower_name(root)} market by region, {Y0 - 1}A-{Y1}F "
                                   f"(USD {'billion' if g['v33'] >= 1000 else 'million'})")
        elif i == 67:
            b.image(i, charts["regions_stacked"])
        elif i == 70:
            b.caption(i, "Figure", f"Global {fmt.lower_name(root)} market scenarios: conservative, base and optimistic")
        elif i == 71:
            b.image(i, charts["scenarios"])
        elif i == 73:
            b.caption(i, "Table", "Scenario definitions")
        elif i == 74:
            cond = obj("scenario_conditions")
            drv = (m.get("fleet_drivers") or {}).get("drivers") or []
            rst = (m.get("fleet_drivers") or {}).get("restraints") or []
            d0 = max(drv, key=lambda x: x.get("cagr_contribution_bps", 0))["name"].lower() if drv else "core demand"
            r0 = max(rst, key=lambda x: x.get("cagr_drag_bps", 0))["name"].lower() if rst else "structural headwinds"
            sdef = {"Conservative": f"Restraints such as {r0} weigh more heavily than in the base case.",
                    "Base": f"Current drivers, led by {d0}, continue at their present pace.",
                    "Optimistic": f"Faster momentum from {d0}, with easing pressure from {r0}."}
            rows = [[n, fmt.pct(v["cagr"]), fmt.usd(v["series"][Y1]), cond.get(n) or sdef[n]]
                    for n, v in m["scenarios"].items()]
            b.table(i, rows)
        elif i == 81:
            b.para(i, f"3.2 Global {root} Market Size by Segments")
        elif i == 83:
            b.para(i, f"Global Opportunity by {tdim['singular']}")
        elif i == 84:
            b.caption(i, "Figure", f"Global market by {tl}, {Y0}E vs {Y1}F")
        elif i == 85:
            b.image(i, charts["type_Global"])
        elif i == 86:
            b.keep(i)        # the allocation rules are listed once, in the methodology's scope table
        elif i == 87:
            b.para(i, seg_para("Global", "type"))
            region_mix(b, "type", tdim["singular"])
        elif 387 <= i <= 392:
            # 'What we do' and 'How we work' are left out: About CMI ends with its opening paragraph
            # (user decision 2026-10-17)
            continue
        elif i == 386:
            el = b.el(i)
            lead = "Coherent Market Insights (CMI)"
            rest = b.src[i].xpath("string(.)").strip()[len(lead):].strip()
            b.out.append(set_lead(el, lead, f"{rest} This report was prepared by CMI's {practice}."))
        elif i == 88:
            b.para(i, f"Global Opportunity by {ddim['singular']}")
        elif i == 89:
            b.caption(i, "Figure", f"Global market by {dl}, {Y0}E vs {Y1}F (USD)")
        elif i == 90:
            b.image(i, charts["demand_Global"])
        elif i == 92:
            b.para(i, seg_para("Global", "demand"))
            region_mix(b, "demand", ddim["singular"])
        elif i == 94:
            b.para(i, f"3.3 Global {root} Market Size by Region")
        elif i == 95:
            b.caption(i, "Figure", f"Regional share of the global market, {Y0 - 1}A")
        elif i == 97:
            b.para(i, s("share_para"))
        elif i == 98:
            b.image(i, charts["region_donuts"])
        elif i == 100:
            b.para(i, f"Global: Absolute dollar opportunity by region, {Y0}-{Y1}", True)
        elif i == 102:
            b.para(i, s("abs_para"))
        elif i == 103:
            b.caption(i, "Figure", f"Global: Absolute dollar opportunity by region, {Y0}-{Y1}")
        elif i == 104:
            b.image(i, charts["region_opportunity"])
        elif i == 107:
            # two regions only: the largest market and the highest CAGR (user decision 2026-10-18)
            hl = obj("highlights")
            for label, r in highlight_regions(m):
                d = m["regions"][r]
                text = re.sub(rf"^\s*{re.escape(r)}[.:]\s*", "", hl.get(r) or "") or (
                    f"The largest regional market, at {fmt.usd(d['v26'])} in {Y0}E." if label.startswith("Largest")
                    else f"The fastest-growing region, at a {fmt.pct(d['cagr'])} CAGR to {Y1}F.")
                b.lead(i, f"{label}: {r}.", text)
        elif i in (108, 109, 110, 111):
            continue
        elif i == 112:
            b.para(i, s("country_bridge"))
        elif any(i in t for t in COUNTRY_TABLES.values()):
            region = next(r for r, t in COUNTRY_TABLES.items() if i in t)
            head, intro, tbl = COUNTRY_TABLES[region]
            if i == head:
                b.para(i, f"{region} {root} Market by Country")
            elif i == intro:
                b.para(i, obj("country_intro").get(region))
            else:
                b.table(i, country_rows(region)[0], header=COUNTRY_HEADER)
        elif any(i == t[2] + 1 for t in COUNTRY_TABLES.values()):
            # the table's source line, then a short profile of every country in the table
            region = next(r for r, t in COUNTRY_TABLES.items() if i == t[2] + 1)
            b.keep(i)
            if region in m["profiled"]:
                continue                        # full profiles with segment charts follow in Chapters 4/5
            profiles = obj("country_profiles")
            why = m.get("geography_why") or {}
            for c, v in m["countries"][region].items():
                # the table above carries the numbers; the profile explains them
                text = with_lead(c, profiles.get(c) or safe(why.get(c)), v, m["regions"][region], region,
                                 m["countries"][region])
                b.lead(29, f"{c}:", text.strip())
        # ------------------------------------------------------------ chapters 4 and 5
        elif i in (146, 226):
            sub = "Middle East" if i == 146 else "Africa"
            b.out.append(set_h1(b.el(i), title=f"{sub} {root} Market: Opportunity Landscape"))
        elif i in (147, 227):
            sub, ch = ("Middle East", "ch4") if i == 147 else ("Africa", "ch5")
            d = m["subregions"][sub]
            lead_in = (f"The {sub} {fmt.lower_name(root)} market is estimated at {fmt.usd(d['v26'])} in {Y0}E and is "
                       f"forecast to reach {fmt.usd(d['v33'])} by {Y1}F ({fmt.pct(d['cagr'])} CAGR).")
            other = "Africa" if sub == "Middle East" else "Middle East"
            # opens on its own sub-region by name (not 'South Africa'), not on the other sub-region
            b.para(i, open_chapter(s(f"{ch}.region_para"),
                                   rf"(?<!South )(?<!Rest of )\b{re.escape(sub)}\b(?! &)", lead_in,
                                   avoid=rf"(?<!South )(?<!Rest of )\b{re.escape(other)}\b(?! &)"))
        elif i in (148, 228):
            sub = "Middle East" if i == 148 else "Africa"
            unit = "billion" if m["subregions"][sub]["v33"] >= 1000 else "million"
            b.caption(i, "Figure", f"{sub} {fmt.lower_name(root)} market by country, {Y0 - 1}A-{Y1}F (USD {unit})")
        elif i in (149, 229):
            b.image(i, charts["me_stacked" if i == 149 else "af_stacked"])
        elif i in (152, 232):
            b.caption(i, "Table", f"{'Middle East' if i == 152 else 'Africa'} {fmt.lower_name(root)} market by country (USD)")
        elif i in (153, 233):
            rows, total = country_rows("Middle East" if i == 153 else "Africa")
            b.table(i, rows, total, header=COUNTRY_HEADER)
        elif i in (155, 235):
            b.caption(i, "Figure", "Country attractiveness: growth vs absolute opportunity")
        elif i in (156, 236):
            b.image(i, charts["me_bubble" if i == 156 else "af_bubble"])
        elif i in (158, 238):
            b.para(i, f"{'4.2 Middle East' if i == 158 else '5.2 Africa'} Opportunity by {tl}")
        elif i in (159, 239):
            b.caption(i, "Figure", f"{'Middle East' if i == 159 else 'Africa'} market by {tl}, {Y0}E vs {Y1}F")
        elif i in (160, 240):
            b.image(i, charts["type_Middle East" if i == 160 else "type_Africa"])
        elif i in (162, 242):
            b.para(i, seg_para("Middle East" if i == 162 else "Africa", "type"))
        elif i in (163, 243):
            b.para(i, f"{'4.3 Middle East' if i == 163 else '5.3 Africa'} Opportunity by {dl}")
        elif i in (164, 244):
            b.caption(i, "Figure", f"{'Middle East' if i == 164 else 'Africa'} market by {dl}, {Y0}E vs {Y1}F (USD)")
        elif i in (165, 245):
            b.image(i, charts["demand_Middle East" if i == 165 else "demand_Africa"])
        elif i in (167, 247):
            b.para(i, seg_para("Middle East" if i == 167 else "Africa", "demand"))
        elif i in (168, 248):
            b.para(i, f"{'4.4 Middle East' if i == 168 else '5.4 Africa'} {root} Market: Country landscape")
        elif i in PROFILE_STARTS:
            # one full profile per named country + the 'Rest of' group, each cloned from the
            # template's first profile block (12 elements); the other template blocks are dropped
            sub = PROFILE_STARTS[i]
            for k, unit in enumerate(m["profiled"][sub]):
                if k and sub == "Middle East":
                    b.keep(181)                                     # page break between profiles
                d = m["countries"][sub][unit]
                for off in range(12):
                    j = i + off
                    if off in range(2, 7) and not sec["type"]["show"] or off in range(7, 12) and not sec["demand"]["show"]:
                        continue
                    if off == 0:
                        b.para(j, unit)
                    elif off == 1:
                        # a profile always opens with its lead sentence (size, share, CAGR)
                        b.para(j, with_lead(unit, s(f"profile__{unit}"), d, m["subregions"][sub],
                                            "the Middle East" if sub == "Middle East" else sub, m["countries"][sub]))
                    elif off == 2:
                        b.para(j, f"{unit} Opportunity by {tl}")
                    elif off == 3:
                        b.caption(j, "Figure", f"{unit} market by {tl}, {Y0}E vs {Y1}F")
                    elif off == 4:
                        b.image(j, charts[f"type_{unit}"])
                    elif off == 6:
                        b.para(j, seg_para(unit, "type"))
                    elif off == 7:
                        b.para(j, f"{unit} Opportunity by {dl}")
                    elif off == 8:
                        b.caption(j, "Figure", f"{unit} market by {dl}, {Y0}E vs {Y1}F (USD)")
                    elif off == 9:
                        b.image(j, charts[f"demand_{unit}"])
                    elif off == 11:
                        b.para(j, seg_para(unit, "demand"))
                    else:
                        b.keep(j)
        elif i in PROFILE_RANGE:
            continue                                                # replaced by the cloned blocks above
        elif i in (221, 298):
            b.para(i, f"{'4.5 Middle East' if i == 221 else '5.5 Africa'} {root} Market: Use-case opportunity heat map")
        elif i in (222, 299):
            b.caption(i, "Table", "Use-case attractiveness by country (CMI assessment)")
        elif i in (223, 300):
            key = "middle_east" if i == 223 else "africa"
            # columns = this market's named countries (template slots: 4 + Rest for ME, 5 for Africa)
            cols = [short(c) for c in m["geography"]["Middle East"]] + ["Rest of ME"] if i == 223 else \
                [short(c) for c in m["geography"]["Africa"]]
            rows = [[r.get("use_case", "")] + [r.get(c, "") for c in cols]
                    for r in (ag.get("use_cases") or {}).get(key) or []]
            b.table(i, complete_rows(rows, b.log, "use-case heat map"), header=["Use case"] + cols)
        elif i in (225, 302):
            ch = "ch4" if i == 225 else "ch5"
            items = [(lead.rstrip(":") + ":", text) for lead, text in pairs(f"{ch}.where_to_play")]
            b.box(i, f"Where to play in {'the Middle East' if i == 225 else 'Africa'}", items)
        # ------------------------------------------------------------ chapter 6: 6.1 Global, 6.2 MEA
        elif i == 303:
            b.out.append(set_h1(b.el(i), title="Regulatory & Policy Outlook: Global vs Middle East & Africa"))
        elif i == 304:
            # 6.1 Global (changes.txt 2026-10-19): the main frameworks outside MEA (each confirmed by a web
            # search, extras.global_regulation) and the global policy outlook
            g_rows = [[r[0], r[1], safe(r[2])] for r in greg.get("frameworks") or []]
            g_out = [x for x in (safe(t) for t in greg.get("outlook") or []) if x]
            if g_rows or g_out:
                b.para(311, "6.1 Global")
                b.para(i, safe(greg.get("intro")))
                if g_rows:
                    b.heading3("Regulatory frameworks")
                    b.caption(312, "Table", "Key regulatory frameworks outside the Middle East & Africa")
                    b.table(313, complete_rows(g_rows, b.log, "global regulatory frameworks"),
                            header=["Jurisdiction", "Framework", "What it means for the market"])
                    b.para(314, "Source: Public regulatory publications; CMI analysis")
                if g_out:
                    b.heading3("Policy outlook to 2030")
                    b.bullets(316, g_out)
                # the global watchlist, like the MEA one (CEO feedback 2026-10-20)
                g_watch = complete_rows([[str(r.get("theme") or ""), safe(r.get("what_to_watch")),
                                          safe(r.get("implication"))] for r in ex.get("global_watchlist") or []],
                                        b.log, "global regulatory watchlist")
                if g_watch:
                    b.heading3(f"Regulatory watchlist {Y0}-{Y0 + 2}")
                    b.caption(321, "Table", "Regulatory watchlist: Global")
                    b.table(322, g_watch)
                    b.para(323, "Source: CMI analysis")
                b.para(311, "6.2 Middle East & Africa")
            b.para(i, open_chapter(s("reg_intro"), r"(?i)regulat|polic|approv|registr|licens|authorit",
                                   f"Regulation decides how quickly {fmt.lower_name(root)} products reach buyers in the Middle "
                                   f"East & Africa: which authority approves them, how long registration takes and "
                                   f"what is changing next."))
        elif i == 305:
            b.caption(i, "Figure", f"Regulatory milestones shaping the {fmt.lower_name(root)} market")
        elif i == 306:
            b.image(i, charts["timeline"])
        elif i == 307:
            b.para(i, "Source: Public regulatory notices; CMI analysis")
        elif i == 308:
            b.caption(i, "Figure", "Regulatory enablement scorecard: leading Middle East & Africa countries")
        elif i == 309:
            b.image(i, charts.get("reg_scorecard"))     # None (no confirmed country): figure left out
        elif i == 312:
            if not reg_rows:
                continue                          # nothing confirmed: caption and table left out
            b.caption(i, "Table", "Regulatory snapshot: Middle East & Africa")
        elif i == 313:
            if reg_rows:
                b.table(i, reg_rows)
            else:
                b.log.append("regulatory snapshot left out: no country with a confirmed authority and requirements")
        elif i in (311, 315, 320):
            # the MEA sub-sections are 6.2.x under '6.2 Middle East & Africa' when 6.1 Global is shown
            title = {311: "Regulator and rules snapshot", 315: "Policy outlook to 2030",
                     320: f"Regulatory watchlist {Y0}-{Y0 + 2}"}[i]
            if greg.get("frameworks") or greg.get("outlook"):
                b.heading3(title)
            else:
                b.para(i, f"6.1 {title}")
        elif i == 316:
            b.bullets(i, lst("policy_bullets"))
        elif i in (317, 318, 319):
            continue
        elif i == 321:
            b.caption(i, "Table", "Regulatory watchlist: Middle East & Africa")
        elif i == 322:
            rows = [[r.get("theme", ""), safe(r.get("what_to_watch")), safe(r.get("implication"))]
                    for r in (ag.get("regulation") or {}).get("watchlist") or []]
            b.table(i, complete_rows(rows, b.log, "regulatory watchlist"))
        elif i == 324:
            b.box(i, "Implication for market entry", [s("entry_implication", "")] if s("entry_implication") else [])
        # ------------------------------------------------------------ chapter 7 (global view)
        elif i == 325:
            b.out.append(set_h1(b.el(i), title="Vendor Intelligence: Global"))
        elif i == 326:
            b.para(i, open_chapter(s("vendor_intro"), r"(?i)vendor|compan|player|supplier|tier|producer",
                                   f"The {fmt.lower_name(root)} vendor landscape is set out below by "
                                   + ("market share and capability." if sec["vendor_global"]["show"]
                                      else "market position and capability.")))
        elif i == 328:
            b.caption(i, "Figure", f"{root} value chain and share of value pool")
        elif i == 329:
            b.image(i, charts["value_chain"])
        elif i == 331:
            b.para(i, s("value_chain_para"))
        elif 332 <= i <= 341:
            if i == 332 and charts.get("tiers"):
                # the global landscape by tier, manufacturers only (changes.txt 2026-10-19)
                b.para(332, "7.2 Global market share by vendor tier")
                b.caption(333, "Figure", f"Global {fmt.lower_name(root)} market: share of value by manufacturer "
                                         f"tier, {Y0 - 1}")
                b.image(334, charts["tiers"])
                b.para(335, "Source: CMI analysis")
                b.para(336, tier_text(ex.get("tiers") or []))
            vendor_shares(b, i, m, s, root, charts, vsafe)
        elif i == 343:
            b.para(i, "7.4 Vendor capability mapping")
        elif i == 347:
            b.para(i, "7.5 Vendor dashboard")
        elif i == 344:
            b.caption(i, "Figure", "Capability heat map of leading vendors")
        elif i == 345:
            b.image(i, charts["capabilities"])
        elif i == 348:
            b.caption(i, "Table", "Vendor dashboard")
        elif i == 349:
            # vendors with a confirmed HQ only; HQ and ownership note from the entity check (PLAN-05);
            # vendor moves that could not be confirmed were removed by verify_facts.apply (PLAN-06), so no marker
            rows = []
            for r in m["vendor_dashboard"]:
                own = short_note(r["ownership_note"])
                foot = vsafe(r["mea_footprint"])
                rows.append([r["vendor"] + (f" ({own})" if own and own.lower() not in r["vendor"].lower() else ""),
                             r["hq"], vsafe(r["core_strength"]), foot or ""])
            b.table(i, complete_rows(rows, b.log, "vendor dashboard"))
        elif i == 351:
            b.box(i, "Competitive implications", lst("implications"), numbered=True)
        # ------------------------------------------------------------ chapter 8
        elif i == 353:
            b.para(i, s("reco_intro") or "The recommendations below turn the report's findings into priority moves "
                                         "for each stakeholder group in the Middle East & Africa.")
        elif i in (354, 355, 356):
            continue                          # the MEA phases figure sits in 8.1.2
        elif i == 358:
            gl_rows = [[str(r["stakeholder"]), safe(str(r["priority_moves"])), safe(str(r["kpis"]))]
                       for r in ex.get("global_stakeholders") or []]
            if gl_rows:
                b.heading3("Global")
                b.caption(i, "Table", "Strategic recommendations by stakeholder: Global")
                b.table(359, complete_rows(gl_rows, b.log, "global stakeholder recommendations"))
                b.para(360, "Source: CMI analysis")
                if charts.get("phases_global"):
                    # the global roadmap, above the MEA one (changes.txt 2026-10-19)
                    b.caption(354, "Figure", "Three-phase way forward for global entry and scale-up")
                    b.image(355, charts["phases_global"])
                    b.para(356, "Source: CMI analysis")
                b.heading3("Middle East & Africa")
            b.caption(i, "Table", "Strategic recommendations by stakeholder: Middle East & Africa")
        elif i == 359:
            rows = [[r.get("stakeholder", ""), safe(r.get("priority_moves")), safe(r.get("kpis"))]
                    for r in (ag.get("strategy") or {}).get("stakeholders") or []]
            b.table(i, complete_rows(rows, b.log, "stakeholder recommendations"))
        elif i == 360:
            b.keep(i)
            b.caption(354, "Figure", "Three-phase way forward for MEA entry and scale-up")
            b.image(355, charts["phases"])
            b.para(356, "Source: CMI analysis")
        elif i == 361:
            # B9: each action carries the finding it follows from (written after the body)
            traced = [(lead.rstrip(".") + ".", text) for lead, text in pairs("priority_actions")
                      if text.startswith("Finding")]
            acts = traced or [x for x in (safe(a) for a in (ag.get("strategy") or {}).get("priority_actions") or []) if x]
            acts = acts[:5]
            # the title states how many actions survived the guards ('Three priority actions ...')
            b.box(i, f"{_COUNT.get(len(acts), str(len(acts)))} priority actions for {Y0}-{str(Y0 + 1)[2:]}",
                  acts, numbered=True)
        # ------------------------------------------------------------ chapter 9
        elif i == 363:
            if sources:
                # 9.1 Sources opens the chapter: the sources used for this market, in the VP's format a-h
                # (changes.txt 2026-10-19); the methodology framework follows as 9.2
                b.para(367, "9.1 Sources")
                b.para(i, f"The {fmt.lower_name(root)} market was researched with the primary, secondary and "
                          f"proprietary sources below, covering the global market and the Middle East & Africa. "
                          f"Consulting firms and commercial market-research publishers are not used as sources.")
                for k, (label, items) in enumerate(sources):
                    b.lead(50, f"{chr(ord('a') + k)}. {label}:", "; ".join(items) + ".")
                b.para(367, "9.2 Research approach")
            # exactly what the pipeline does (PLAN-07): no expert panel, interviews, bottom-up or triangulation
            b.para(i, method_text(m, root))
        elif i == 364:
            b.caption(i, "Figure", "CMI's Integrated Market Research Methodology Framework")
        elif i == 365:
            b.image(i, charts["process"])
        elif i in (367, 368):
            # no 'basis of the headline' section: where the headline comes from is internal
            # (market_size_check.json) and the client report makes no claims about it (2026-10-16)
            continue
        elif i in (369, 370, 371):
            pub = []                 # third-party publishers are never named in the client report (user decision)
            if not pub:
                continue
            if i == 369:
                b.caption(i, "Figure", f"Published {Y0 - 1} global {fmt.lower_name(root)} market estimates vs CMI base")
            elif i == 370:
                b.image(i, charts["published"])
            else:
                b.para(i, "Source: " + ", ".join(dict.fromkeys(p["publisher"] for p in pub)) + "; CMI analysis")
        elif i in (372, 373, 374, 375):
            # reference figures only beside a sourced headline; they never set the CMI numbers
            refs = []                # third-party publishers are never named in the client report (user decision)
            if not refs:
                continue
            if i == 373:
                b.caption(i, "Table", "Third-party reference figures (for reference only; not used in the model)")
            elif i == 374:
                b.table(i, [[r["name"], r.get("scope", ""), r["published_value"]] for r in refs])
            else:
                b.keep(i)
        elif i == 377:
            b.caption(i, "Table", "Scope of the study")
        elif i == 378:
            dfn = m["definition"]
            rows = [["Market measured", f"{dfn.get('measures') or dfn['definition']} ({dfn['unit']})"]] \
                if dfn.get("measures") or dfn.get("definition") else []
            if dfn.get("value_chain_level"):
                rows.append(["Value-chain level", dfn["value_chain_level"]])
            rows += [[(tdim if d == "type" else ddim)["singular"], "; ".join(seg["Global"][d])]
                     for d in ("type", "demand") if sec[d]["show"]]
            if dfn.get("included"):
                rows.append(["Inclusions", dfn["included"]])
            if dfn.get("excluded"):
                rows.append(["Exclusions", dfn["excluded"]])
            # exactly the countries covered in the body
            gl = []
            for r in ("North America", "Europe", "Asia Pacific", "Latin America", "Middle East", "Africa"):
                named = [c for c in m["countries"][r] if not c.startswith("Rest of")]
                rest = [c for c in m["countries"][r] if c.startswith("Rest of")]
                gl.append(f"{r} ({', '.join(named + rest)})")
            rows.append(["Geography", "; ".join(gl)])
            for d in ("type", "demand"):
                for n, txt in (m.get("segment_definitions") or {}).get(d, {}).items():
                    if sec[d]["show"] and n in seg["Global"][d]:
                        rows.append([f"{(tdim if d == 'type' else ddim)['singular']}: {n}", txt])
            for rule in m.get("allocation_rules") or []:
                rows.append(["Allocation rule", rule])
            rows.append(["Years", f"Base {Y0 - 1}; estimate {Y0}; forecast {Y0 + 1}-{Y1}"])
            b.table(i, rows)
        elif i in (380, 381, 382, 383, 384):
            continue             # 'Assumptions and limitations' reads as a disclaimer: left out (2026-10-16)
        # ------------------------------------------------------------ captions/sources kept as-is
        else:
            text = b.src[i].xpath("string(.)").strip()
            m_cap = re.match(r"^(Figure|Table) \d+:", text)
            if m_cap:
                b.caption(i, m_cap.group(1), text.split(":", 1)[1].strip())
            elif re.match(r"Source: CMI analysis\. \d{4}E = estimate", text):
                b.para(i, f"Source: CMI analysis. A = Actual ({Y0 - 1}), E = Estimated ({Y0}), "
                          f"F = Forecast ({Y0 + 1}-{Y1}).")
            else:
                b.keep(i)

    # no client appendix (master data / registers): the charts and tables carry the numbers, and the full
    # number register is the CSV beside the report (user decision 2026-10-16: a clean client document)
    b.drop_empty_headings()
    if internal is not None:
        internal.update(internal_log(b, m, ag, prose, removed_notes, verification))
    renumber(b, toc_titles, toc_pages)
    b.save(out_path, f"The {root} Market Outlook {Y0}-{Y1}", "")      # no "Whitepaper" footer label (2026-10-20)
    return removed_notes + b.log + [f"guard dropped: {d}" for d in dropped]
